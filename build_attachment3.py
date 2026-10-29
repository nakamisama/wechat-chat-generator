# -*- coding: utf-8 -*-
"""把聊天记录截图按模板格式写入「附件3.项目进度沟通记录」Word（纯截图，每张一页）。

项目路径与截图前缀放在本地 local/attachment3.json，脚本本身不含客户信息。

用法：
    python build_attachment3.py
    python build_attachment3.py --data local/其他.json
"""
import argparse, json, os
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_DATA = os.path.join(BASE, 'local', 'attachment3.json')


def build_one(doc_path, shot_dir, prefix):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)

    files = sorted([f for f in os.listdir(shot_dir)
                    if f.startswith(prefix) and f.endswith('.jpg')])

    for i, f in enumerate(files):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_path = os.path.join(shot_dir, f)
        with Image.open(img_path) as im:
            iw, ih = im.size
        if iw / ih < 1:
            # 竖屏（手机截图）：按高度适配
            p.add_run().add_picture(img_path, height=Cm(24.0))
        else:
            # 横屏（电脑截图）：按宽度适配
            p.add_run().add_picture(img_path, width=Cm(14.6))
        if i < len(files) - 1:
            doc.add_page_break()

    doc.save(doc_path)
    print(f'  OK: {doc_path}  ({len(files)} 张)')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data', default=DEFAULT_DATA, help='项目清单 JSON（本地）')
    args = ap.parse_args()

    if not os.path.exists(args.data):
        print(f'数据文件不存在: {args.data}')
        print('提示：该文件含客户信息，只保留本地，换电脑请用 project_config.py 迁移')
        return

    data = json.load(open(args.data, encoding='utf-8'))
    for job in data['jobs']:
        shot_dir = os.path.join(job['dir'], '聊天记录截图')
        if not os.path.isdir(shot_dir):
            print(f'  跳过（无截图目录）: {shot_dir}')
            continue
        doc_path = os.path.join(job['dir'], job.get('doc_name', '附件3.项目进度沟通记录.docx'))
        build_one(doc_path, shot_dir, job['prefix'])


main()
