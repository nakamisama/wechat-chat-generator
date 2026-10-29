# -*- coding: utf-8 -*-
"""生成「附件7.送货单」Word 文档（金额栏留空手写）。

项目数据（收货单位、单号、日期、物料明细）放在本地 local/delivery_notes.json，
脚本本身不含任何客户信息。

用法：
    python gen_delivery_notes.py
    python gen_delivery_notes.py --data local/其他.json
"""
import argparse, json, os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_DATA = os.path.join(BASE, 'local', 'delivery_notes.json')


def set_cell_font(cell, size=10, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.bold = bold
            r.font.name = '仿宋'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def build(items, out_path, title, buyer, seller, delivery_no, delivery_date, extra_note=''):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = h.add_run(title)
    hr.bold = True
    hr.font.size = Pt(20)
    hr.font.name = '黑体'
    hr._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    doc.add_paragraph()

    info = [
        ('收货单位：', buyer),
        ('送货单位：', seller),
        ('送货单号：', delivery_no),
        ('送货日期：', delivery_date),
    ]
    for label, val in info:
        p = doc.add_paragraph()
        for txt in (label, val):
            r = p.add_run(txt)
            r.font.size = Pt(11)
            r.font.name = '仿宋'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

    doc.add_paragraph()

    header = ['序号', '名称', '规格型号', '单位', '数量', '单价', '合计', '备注']
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for i, t in enumerate(header):
        hdr[i].text = t
        set_cell_font(hdr[i], 10, True)
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, it in enumerate(items, start=1):
        row = table.add_row().cells
        row[0].text = str(idx)
        row[1].text = it[0]
        row[2].text = it[1]
        row[3].text = it[2]
        row[4].text = str(it[3])
        row[5].text = ''   # 单价 留空手写
        row[6].text = ''   # 合计 留空手写
        row[7].text = it[4] if len(it) > 4 else ''
        for c in row:
            set_cell_font(c, 10)
        for i in (0, 3, 4, 5, 6):
            for p in row[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    total = table.add_row().cells
    for c in total:
        c.text = ''
    merged = total[0].merge(total[1]).merge(total[2])
    merged.text = '合计'
    for p_ in merged.paragraphs:
        p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for c in total:
        set_cell_font(c, 10, True)

    set_col_widths(table, [1.3, 4.6, 3.6, 1.2, 1.2, 1.8, 1.8, 2.1])

    doc.add_paragraph()

    if extra_note:
        pn = doc.add_paragraph()
        rn = pn.add_run(extra_note)
        rn.font.size = Pt(9)
        rn.font.name = '仿宋'
        rn._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
        doc.add_paragraph()

    for line in [
        '收货人（签字）：________________          联系电话：________________',
        '送货人（签字）：________________          联系电话：________________',
        '签收日期：________年______月______日',
    ]:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(11)
        r.font.name = '仿宋'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')
        p.paragraph_format.space_after = Pt(14)

    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run('说明：本单一式两联，收货方、送货方各执一联。')
    r.font.size = Pt(9)
    r.font.name = '仿宋'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋')

    doc.save(out_path)
    print(f'  已生成 {out_path}  ({len(items)} 项)')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data', default=DEFAULT_DATA, help='送货单数据 JSON（本地）')
    ap.add_argument('--seller', default=None, help='送货单位，默认取数据文件里的 seller')
    args = ap.parse_args()

    if not os.path.exists(args.data):
        print(f'数据文件不存在: {args.data}')
        print('提示：该文件含客户信息，只保留本地，换电脑请用 project_config.py 迁移')
        return

    data = json.load(open(args.data, encoding='utf-8'))
    seller = args.seller or data.get('seller', '')

    for job in data['jobs']:
        out_dir = job['out_dir']
        os.makedirs(out_dir, exist_ok=True)
        out = os.path.join(out_dir, job.get('doc_name', '附件7.送货单.docx'))
        build(
            job['items'], out,
            job.get('title', '送 货 清 单'),
            job['buyer'], seller,
            job['no'], job['date'],
            job.get('note', ''),
        )


main()
